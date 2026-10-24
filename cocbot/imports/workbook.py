"""Open an uploaded .xlsx safely, read-only, without executing anything.

Checks happen on the ZIP container *before* openpyxl sees it, so oversized
or malformed uploads are refused cheaply. The workbook is loaded twice from
memory: once with formulas and once with Excel's cached values. Neither load
evaluates formulas, runs macros or follows external links. Nothing is
written to disk.
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import Any, List, Optional

from cocbot.i18n import Msg

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ENTRIES = 200
MAX_TOTAL_UNCOMPRESSED = 25 * 1024 * 1024
MAX_ENTRY_UNCOMPRESSED = 10 * 1024 * 1024
MAX_RATIO = 100
MAX_ROWS, MAX_COLS = 500, 40
REQUIRED_PARTS = ("[Content_Types].xml", "xl/workbook.xml")


class WorkbookRejected(Exception):
    """The upload cannot be read safely. ``args[0]`` is a user-facing Msg (never file contents)."""


@dataclass
class Book:
    """Formula view and cached-value view of the same workbook."""
    formulas: Any            # openpyxl Workbook (data_only=False)
    values: Any              # openpyxl Workbook (data_only=True)
    notices: List[Msg] = field(default_factory=list)   # e.g. external links ignored

    @property
    def sheet_names(self) -> List[str]:
        return list(self.formulas.sheetnames)

    def has(self, sheet: str) -> bool:
        return sheet in self.formulas.sheetnames

    def raw(self, sheet: str, ref: str) -> Any:
        """The stored cell content: a number, a string, or a formula string starting with '='."""
        return self.formulas[sheet][ref].value

    def cached(self, sheet: str, ref: str) -> Any:
        """Excel's last computed value (None if the file was saved without one)."""
        return self.values[sheet][ref].value


def check_container(data: bytes) -> List[Msg]:
    """Validate the ZIP structure; returns notices, raises WorkbookRejected."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise WorkbookRejected(Msg("import_too_large", mb=MAX_UPLOAD_BYTES // (1024 * 1024)))
    if not data.startswith(b"PK"):
        raise WorkbookRejected(Msg("import_not_xlsx"))
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        raise WorkbookRejected(Msg("import_not_xlsx")) from None
    if len(infos) > MAX_ENTRIES:
        raise WorkbookRejected(Msg("import_too_complex"))
    total = 0
    names = set()
    for info in infos:
        names.add(info.filename)
        if info.flag_bits & 0x1:                          # encrypted entry
            raise WorkbookRejected(Msg("import_encrypted"))
        if info.file_size > MAX_ENTRY_UNCOMPRESSED:
            raise WorkbookRejected(Msg("import_too_complex"))
        if info.compress_size and info.file_size / info.compress_size > MAX_RATIO:
            raise WorkbookRejected(Msg("import_too_complex"))
        total += info.file_size
    if total > MAX_TOTAL_UNCOMPRESSED:
        raise WorkbookRejected(Msg("import_too_complex"))
    if any(part not in names for part in REQUIRED_PARTS):
        raise WorkbookRejected(Msg("import_not_xlsx"))
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise WorkbookRejected(Msg("import_macros"))
    notices: List[Msg] = []
    if any(n.startswith("xl/externalLinks/") for n in names):
        notices.append(Msg("import_external_links"))
    return notices


def open_book(data: bytes) -> Book:
    """Open the upload after container checks. Raises WorkbookRejected."""
    notices = check_container(data)
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - dependency is in requirements.txt
        raise WorkbookRejected(Msg("import_unavailable")) from None
    try:
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False, keep_links=False)
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_vba=False, keep_links=False)
    except Exception:   # openpyxl raises many types on malformed XML; message stays generic
        raise WorkbookRejected(Msg("import_unreadable")) from None
    for ws in formulas.worksheets:
        if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS:
            raise WorkbookRejected(Msg("import_too_complex"))
    return Book(formulas, values, notices)


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def as_int(value: Any) -> Optional[int]:
    """Integer from a numeric cell or a numeric string; None for blanks and text."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if value != value or value in (float("inf"), float("-inf")):   # NaN / inf
            return None
        return int(value) if float(value).is_integer() else int(value // 1)
    if isinstance(value, str):
        s = value.strip().replace(",", "")
        if s.lstrip("+-").isdigit():
            return int(s)
    return None


def as_text(value: Any, limit: int = 200) -> Optional[str]:
    if value is None or is_formula(value):
        return None
    s = " ".join(str(value).split())
    return s[:limit] or None
