"""Regenerate data/occupations/occupations.json from the development spreadsheets.

Developer tool only (needs ``pip install -r requirements-import.txt``). The bot
reads the committed JSON and never needs Excel.

    python scripts/import_occupations.py            # writes JSON + prints the review report
    python scripts/import_occupations.py --check    # exit 1 if the committed JSON is stale

Sources:
* data/seed/COC七版人物卡v1.35.xlsx, sheet 职业列表 (authoritative: Chinese name,
  Credit Rating, skill-point formula, occupation skills)
* data/seed/career.xlsx (English names, matched by row number; formulas are
  only cross-checked)

Only structured game data is extracted (names, numbers, skill structure); no
descriptive text is reproduced.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from cocbot.occupations.zh_parser import (  # noqa: E402
    formula_from_career_xlsx,
    parse_formula_zh,
    parse_skill_text,
)

WORKBOOK = ROOT / "data" / "seed" / "COC七版人物卡v1.35.xlsx"
CAREERS = ROOT / "data" / "seed" / "career.xlsx"
OUT = ROOT / "data" / "occupations" / "occupations.json"

# Corrections to career.xlsx English names (typos, duplicates, mistranslations).
EN_NAME_FIXES = {
    7: "Psychiatrist (Classic)",
    10: "Antique Dealer",
    42: "Amateur Art/Craft Enthusiast",
    52: "Entertainer",
    56: "Firefighter",
    57: "Foreign Correspondent",
}
ERA_TAGS = {"古典": "classic", "现代": "modern"}


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def build() -> tuple[dict, list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    rows = list(wb["职业列表"].iter_rows(values_only=True))[1:]
    careers = {r[0]: r for r in openpyxl.load_workbook(CAREERS, read_only=True, data_only=True)
               .worksheets[0].iter_rows(values_only=True) if isinstance(r[0], int)}

    occupations, report, seen = [], [], set()
    for row in rows:
        no, name_zh, cr, formula_zh, _points, skills_zh = (list(row) + [None] * 6)[:6]
        if not isinstance(no, int) or no <= 1:      # 0 = clear template, 1 = custom (handled as homebrew)
            continue
        name_zh = str(name_zh).strip()
        issues = []

        tag = re.search(r"[（(]([^）)]+)[）)]", name_zh)
        era = ERA_TAGS.get(tag.group(1)) if tag else None
        rulebook_variant = bool(tag and tag.group(1) == "原作向")

        en = EN_NAME_FIXES.get(no) or str(careers.get(no, [None, None])[1] or "").strip()
        if not en:
            issues.append("no English name in career.xlsx")
            en = name_zh
        if era == "classic" and "(Classic)" not in en:
            en += " (Classic)"
        if era == "modern":
            en += " (Modern)"
        key = _slug(en)
        if key in seen:
            raise SystemExit(f"duplicate key {key} (row {no})")
        seen.add(key)

        m = re.fullmatch(r"\s*(\d+)\s*-\s*(\d+)\s*", str(cr or ""))
        if not m:
            issues.append(f"unparseable Credit Rating {cr!r}")
            continue
        cr_min, cr_max = int(m.group(1)), int(m.group(2))

        formula = parse_formula_zh(formula_zh)
        other = formula_from_career_xlsx(careers.get(no, [None] * 5)[4])
        if other is not None and other != formula:
            issues.append(f"formula conflict: workbook {formula_zh} vs career.xlsx {careers[no][4]}; workbook used")

        parsed = parse_skill_text(skills_zh)
        issues += parsed.issues
        occupations.append({
            "key": key,
            "name_en": en,
            "name_zh": name_zh,
            "era": era,
            "rulebook_variant": rulebook_variant,
            "source_row": no,
            "credit_rating": [cr_min, cr_max],
            "skill_points": formula,
            "slots": [s.as_dict() for s in parsed.slots],
            "keeper_note_zh": parsed.keeper_note_zh,
            "verified": False,
            "issues": issues,
        })
        if issues:
            report.append(f"#{no} {name_zh} / {en}: " + "; ".join(issues))

    data = {
        "schema": 1,
        "source": {
            "workbook": WORKBOOK.name,
            "sheet": "职业列表",
            "english_names": CAREERS.name,
            "note": "Structured data (names, Credit Rating ranges, skill-point formulas, occupation skill "
                    "structure) extracted from a community character sheet for Call of Cthulhu 7th Edition. "
                    "Call of Cthulhu is a trademark of Chaosium Inc. No rulebook text is reproduced. "
                    "Entries are unverified against the rulebooks unless 'verified' is true.",
        },
        "occupations": occupations,
    }
    return data, report


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    data, report = build()
    text = json.dumps(data, ensure_ascii=False, indent=1) + "\n"
    if args.check:
        current = OUT.read_text(encoding="utf-8") if OUT.exists() else ""
        if current != text:
            print("occupations.json is stale; rerun scripts/import_occupations.py")
            raise SystemExit(1)
        print("occupations.json is up to date")
        return
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(text, encoding="utf-8")
    print(f"wrote {len(data['occupations'])} occupations to {OUT.relative_to(ROOT)}")
    print(f"{len(report)} need review:")
    for line in report:
        print("  " + line)


if __name__ == "__main__":
    main()
