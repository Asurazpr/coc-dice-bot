"""Synthetic .xlsx workbooks in the v1.35 character-sheet layout, built in code.

The real template is third-party (licensing unresolved), so tests never read it.
Only the labels the adapter checks and the cells it reads are reproduced.
``cached`` injects Excel-style saved values into formula cells (openpyxl itself
writes formulas without results, like many non-Excel tools).
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Dict, Mapping, Optional, Tuple

import openpyxl

SHEET = "人物卡"
LABELS = {"B3": "姓名", "I3": "力量\nSTR", "L3": "敏捷\nDEX", "O3": "意志\nPOW", "I5": "体质\nCON",
          "L5": "外表\nAPP", "O5": "教育\nEDU", "I7": "体型\nSIZ", "L7": "智力INT\n灵感idea", "O7": "移动力\nMOV",
          "J10": "幸运\nLuck", "B14": "技能名称", "I14": "成功率", "L14": "技能名称", "S14": "成功率",
          "B4": "玩家", "E4": "时代", "B5": "职业", "B6": "年龄", "E6": "性别", "B7": "住地", "B8": "故乡"}
CHAR_CELLS = {"STR": "J3", "DEX": "M3", "POW": "P3", "CON": "J5", "APP": "M5", "EDU": "P5", "SIZ": "J7",
              "INT": "M7"}

# left block (B name, C spec, E init, F growth, G occ, H interest, I total)
LEFT = {15: "会计", 16: "人类学", 17: "估价", 18: "考古学", 19: "技艺", 22: "魅惑", 23: "攀爬",
        24: '=IF(F4="现代","电脑使用","[不可用]")', 25: "信誉", 26: "克苏鲁神话", 27: "乔装", 28: "闪避",
        29: "汽车驾驶", 30: "电气维修", 31: '=IF(F4="现代","电子学","[不可用]")', 32: "话术", 33: "格斗",
        36: "射击", 39: "急救", 40: "历史", 41: "恐吓", 42: "跳跃", 43: "外语", 46: "母语"}
LEFT_INIT = {15: 5, 16: 1, 17: 5, 18: 1, 19: 5, 20: 5, 21: 5, 22: 15, 23: 20, 24: 5, 25: 0, 26: 0, 27: 5,
             28: "=INT(M3/2)", 29: 20, 30: 10, 31: 1, 32: 5, 33: 25, 34: 0, 35: 0, 36: 20, 37: 0, 38: 0, 39: 30,
             40: 5, 41: 15, 42: 20, 43: 1, 44: 1, 45: 1, 46: "=P5"}
RIGHT = {15: "法律", 16: "图书馆利用", 17: "聆听", 18: "锁匠", 19: "机械维修", 20: "医学", 21: "自然学",
         22: "领航", 23: "神秘学", 24: "操作重型机械", 25: "说服", 26: "驾驶", 27: "精神分析", 28: "心理学",
         29: "骑术", 30: "科学", 33: "妙手", 34: "侦察", 35: "潜行", 36: "生存", 37: "游泳", 38: "投掷",
         39: "追踪", 40: "罕见"}
RIGHT_INIT = {15: 5, 16: 20, 17: 20, 18: 1, 19: 10, 20: 1, 21: 10, 22: 10, 23: 5, 24: 1, 25: 10, 26: 1, 27: 1,
              28: 10, 29: 5, 30: 1, 31: 1, 32: 1, 33: 10, 34: 25, 35: 20, 36: 10, 37: 20, 38: 20, 39: 10, 40: 5}

DEFAULT_CHARS = {"STR": 60, "DEX": 70, "POW": 45, "CON": 55, "APP": 50, "EDU": 75, "SIZ": 65, "INT": 80}


def build(name: Optional[str] = "Ada Lovell", chars: Optional[Mapping[str, object]] = None, luck: object = 50,
          hp: object = None, mp: object = None, san: object = None, occupation: Optional[str] = "会计师",
          occupation_index: Optional[int] = 2, age: object = 32, era: str = "现代",
          points: Optional[Dict[Tuple[str, int], Tuple[int, int, int]]] = None,
          specs: Optional[Dict[Tuple[str, int], str]] = None,
          overrides: Optional[Dict[str, object]] = None,
          extra_sheets: bool = True, signature: bool = True) -> openpyxl.Workbook:
    """``points[(block, row)] = (growth, occupation, interest)``; block 'L' or 'R'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    if signature:
        for ref, text in LABELS.items():
            ws[ref] = text
    ws["C3"] = name
    ws["C5"] = occupation
    ws["C6"] = age
    ws["F4"] = era
    ws["F6"] = "女"
    ws["C4"] = "Player One"
    ws["C7"] = "Arkham"
    ws["C8"] = "Boston"
    for stat, ref in CHAR_CELLS.items():
        ws[ref] = (chars or DEFAULT_CHARS).get(stat, DEFAULT_CHARS[stat]) if chars is None or stat in chars else None
    ws["P7"] = "=8"
    ws["L10"] = luck
    ws["D10"], ws["P10"], ws["H10"] = hp, mp, san
    ws["E10"] = '="/"&INT((J7+J5)/10)'
    ws["Q10"] = '="/"&INT(P3/5)'
    ws["I10"] = '="/"&MIN(P3,99-I26)'
    ws["D12"] = occupation_index
    for row in range(15, 47):
        if row in LEFT:
            ws[f"B{row}"] = LEFT[row]
        if row in LEFT_INIT:
            ws[f"E{row}"] = LEFT_INIT[row]
            ws[f"I{row}"] = f"=SUM(E{row}:H{row})"
        if row in RIGHT:
            ws[f"L{row}"] = RIGHT[row]
        if row in RIGHT_INIT:
            ws[f"O{row}"] = RIGHT_INIT[row]
            ws[f"S{row}"] = f"=SUM(O{row}:R{row})"
    ws["C33"] = "斗殴"
    ws["C36"] = "手枪"
    ws["F25"] = ws["H25"] = ws["G26"] = ws["H26"] = "——"
    for (block, row), spec in (specs or {}).items():
        ws[f"{'C' if block == 'L' else 'M'}{row}"] = spec
        if block == "R" and row > 40:
            ws[f"L{row}"] = spec
            ws[f"S{row}"] = f"=SUM(O{row}:R{row})"
    for (block, row), (g, o, i) in (points or {}).items():
        cols = ("F", "G", "H") if block == "L" else ("P", "Q", "R")
        for col, v in zip(cols, (g, o, i)):
            if v:
                ws[f"{col}{row}"] = v
    ws["M58"] = "Science explains everything"
    ws["K58"] = "信念/信仰"
    ws["B50"], ws["B51"], ws["F51"], ws["I51"] = "空手战斗", ".38 Revolver", 40, "1D10"
    for ref, value in (overrides or {}).items():
        ws[ref] = value
    if extra_sheets:
        occ = wb.create_sheet("职业列表")
        occ.append(["序号", "职业", "信誉", "职业属性", "技能点", "本职技能"])
        occ.append([0, "（无）"])
        occ.append([1, "自定义职业"])
        occ.append([2, "会计师", "30-70", "教育×4", "=人物卡!P5*4", "会计，法律…"])
        wb.create_sheet("分支技能")
    return wb


def to_bytes(wb: openpyxl.Workbook, cached: Optional[Dict[str, object]] = None,
             add_parts: Optional[Dict[str, bytes]] = None) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if not cached and not add_parts:
        return data
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if cached and info.filename == "xl/worksheets/sheet1.xml":
                text = content.decode("utf-8")
                for ref, value in cached.items():
                    text = _inject(text, ref, value)
                content = text.encode("utf-8")
            dst.writestr(info, content)
        for name, content in (add_parts or {}).items():
            dst.writestr(name, content)
    return out.getvalue()


def _inject(xml: str, ref: str, value: object) -> str:
    """Give the formula cell ``ref`` a saved result, as Excel does."""
    pattern = re.compile(rf'<c r="{ref}"([^>]*)><f>(.*?)</f>(<v>.*?</v>|<v\s*/>)?</c>', re.S)
    m = pattern.search(xml)
    if not m:
        raise AssertionError(f"no formula cell {ref}")
    attrs = re.sub(r'\st="[^"]*"', "", m.group(1))
    if isinstance(value, str):
        repl = f'<c r="{ref}"{attrs} t="str"><f>{m.group(2)}</f><v>{value}</v></c>'
    else:
        repl = f'<c r="{ref}"{attrs}><f>{m.group(2)}</f><v>{value}</v></c>'
    return xml[:m.start()] + repl + xml[m.end():]


def sheet_bytes(**kw) -> bytes:
    cached = kw.pop("cached", None)
    add_parts = kw.pop("add_parts", None)
    return to_bytes(build(**kw), cached=cached, add_parts=add_parts)
